from niceexcel_api.main import app, _workbooks

client = app.test_client()


def setup_function():
    _workbooks.clear()
    from niceexcel_api.main import _undo_history, _redo_history
    _undo_history.clear()
    _redo_history.clear()


def test_health_and_create_workbook():
    assert client.get('/health').get_json()['status'] == 'ok'
    response = client.post('/workbooks', json={'sheet_name': 'Budget'})
    assert response.status_code == 200
    data = response.get_json()
    assert data['sheets'] == ['Budget']
    assert data['workbook_id']


def test_cells_sheets_style_and_save_roundtrip():
    book_id = client.post('/workbooks', json={}).get_json()['workbook_id']
    result = client.patch(f'/workbooks/{book_id}/cells', json={
        'sheet': 'Sheet1', 'cells': [
            {'cell': 'A1', 'value': 'Revenue'},
            {'cell': 'B1', 'value': 150},
            {'cell': 'B2', 'value': '=SUM(B1:B1)'},
        ]
    })
    assert result.status_code == 200
    assert client.patch(f'/workbooks/{book_id}/style', json={
        'sheet': 'Sheet1', 'range': 'A1:B1', 'bold': True, 'fill_color': '#D9EAD3'
    }).status_code == 200
    assert client.post(f'/workbooks/{book_id}/sheets', json={'name': 'Forecast'}).status_code == 200
    info = client.get(f'/workbooks/{book_id}').get_json()
    assert [s['name'] for s in info['sheets']] == ['Sheet1', 'Forecast']
    cells = client.get(f'/workbooks/{book_id}/range', query_string={'sheet': 'Sheet1', 'range': 'A1:B2'}).get_json()
    assert cells['rows'][0][0]['value'] == 'Revenue'
    assert cells['rows'][0][0]['style']['bold'] is True
    assert cells['rows'][1][1]['value'] == '=SUM(B1:B1)'
    saved = client.get(f'/workbooks/{book_id}/save')
    assert saved.status_code == 200
    assert saved.data[:2] == b'PK'


def test_insert_row_shifts_existing_cells_and_autosaves():
    book_id = client.post("/workbooks", json={}).get_json()["workbook_id"]
    client.patch(f"/workbooks/{book_id}/cells", json={
        "sheet": "Sheet1", "cells": [
            {"cell": "A1", "value": "Header"},
            {"cell": "A2", "value": "Keep me"},
        ]
    })
    response = client.post(f"/workbooks/{book_id}/rows", json={"sheet": "Sheet1", "row": 2})
    assert response.status_code == 200, response.get_data(as_text=True)
    assert response.get_json()["inserted_row"] == 2
    rows = client.get(f"/workbooks/{book_id}/range", query_string={"sheet": "Sheet1", "range": "A1:A3"}).get_json()["rows"]
    assert rows[0][0]["value"] == "Header"
    assert rows[1][0]["value"] is None
    assert rows[2][0]["value"] == "Keep me"


def test_last_sheet_cannot_be_deleted_and_bad_range_rejected():
    book_id = client.post('/workbooks', json={}).get_json()['workbook_id']
    response = client.delete(f'/workbooks/{book_id}/sheets/Sheet1')
    assert response.status_code == 409
    response = client.get(f'/workbooks/{book_id}/range', query_string={'sheet': 'Sheet1', 'range': 'not-a-range'})
    assert response.status_code == 422


def test_formula_text_is_preserved_and_subset_result_is_calculated():
    book_id = client.post('/workbooks', json={}).get_json()['workbook_id']
    client.patch(f'/workbooks/{book_id}/cells', json={'sheet': 'Sheet1', 'cells': [{'cell': 'A1', 'value': '=SUM(1,2)'}]})
    cell = client.get(f'/workbooks/{book_id}/range', query_string={'sheet': 'Sheet1', 'range': 'A1:A1'}).get_json()['rows'][0][0]
    assert cell['value'] == '=SUM(1,2)'
    assert cell['calculated_value'] == 3.0
    assert client.get(f'/workbooks/{book_id}').get_json()['formula_calculation'] == 'limited_nice_excel_subset_evaluator'


def test_spreadsheet_ui_is_served_from_local_api():
    response = client.get('/')
    assert response.status_code == 200
    assert 'Nice Excel' in response.get_data(as_text=True)
    assert 'formula' in response.get_data(as_text=True).lower()


def test_formula_engine_calculates_basic_functions_and_references():
    book_id = client.post('/workbooks', json={}).get_json()['workbook_id']
    client.patch(f'/workbooks/{book_id}/cells', json={'sheet': 'Sheet1', 'cells': [
        {'cell': 'A1', 'value': 10}, {'cell': 'A2', 'value': 20},
        {'cell': 'B1', 'value': '=SUM(A1:A2)'},
        {'cell': 'B2', 'value': '=IF(A1>5,"Yes","No")'},
        {'cell': 'B3', 'value': '=COUNTIF(A1:A2,">10")'},
        {'cell': 'B4', 'value': '=SUMIF(A1:A2,">10",A1:A2)'},
        {'cell': 'B5', 'value': '=A1+A2*2'},
    ]})
    rows = client.get(f'/workbooks/{book_id}/range', query_string={'sheet': 'Sheet1', 'range': 'B1:B5'}).get_json()['rows']
    assert [row[0]['calculated_value'] for row in rows] == [30.0, 'Yes', 1, 20.0, 50]
    assert rows[0][0]['value'] == '=SUM(A1:A2)'  # formula text is preserved


def test_formula_engine_lookup_and_formula_errors():
    book_id = client.post('/workbooks', json={}).get_json()['workbook_id']
    client.patch(f'/workbooks/{book_id}/cells', json={'sheet': 'Sheet1', 'cells': [
        {'cell': 'A1', 'value': 'Apple'}, {'cell': 'B1', 'value': 12},
        {'cell': 'A2', 'value': 'Pear'}, {'cell': 'B2', 'value': 18},
        {'cell': 'D1', 'value': '=VLOOKUP("Pear",A1:B2,2,FALSE)'},
        {'cell': 'D2', 'value': '=XLOOKUP("Apple",A1:A2,B1:B2,"Missing")'},
        {'cell': 'D3', 'value': '=SUM(A1:A2)'},
        {'cell': 'D4', 'value': '=1/0'},
    ]})
    rows = client.get(f'/workbooks/{book_id}/range', query_string={'sheet': 'Sheet1', 'range': 'D1:D4'}).get_json()['rows']
    assert rows[0][0]['calculated_value'] == 18
    assert rows[1][0]['calculated_value'] == 12
    assert rows[2][0]['calculated_value'] == 0
    assert rows[3][0]['calculated_value'] == '#DIV/0!'


def test_autosave_recovers_workbook_after_session_cache_is_cleared():
    created = client.post('/workbooks', json={'sheet_name': 'Recovery'}).get_json()
    book_id = created['workbook_id']
    changed = client.patch(f'/workbooks/{book_id}/cells', json={
        'sheet': 'Recovery', 'cells': [{'cell': 'A1', 'value': 'Saved automatically'}]
    })
    assert changed.status_code == 200
    assert changed.get_json()['autosaved'] is True
    _workbooks.clear()  # simulate an API process restart/session cache loss
    recovered = client.get(f'/workbooks/{book_id}/range', query_string={
        'sheet': 'Recovery', 'range': 'A1:A1'
    })
    assert recovered.status_code == 200
    assert recovered.get_json()['rows'][0][0]['value'] == 'Saved automatically'


def test_chart_creation_and_recovery_listing():
    created = client.post('/workbooks', json={}).get_json()
    book_id = created['workbook_id']
    client.patch(f'/workbooks/{book_id}/cells', json={'sheet': 'Sheet1', 'cells': [
        {'cell': 'A1', 'value': 'Month'}, {'cell': 'B1', 'value': 'Sales'},
        {'cell': 'A2', 'value': 'Jan'}, {'cell': 'B2', 'value': 10},
        {'cell': 'A3', 'value': 'Feb'}, {'cell': 'B3', 'value': 15},
    ]})
    chart = client.post(f'/workbooks/{book_id}/charts', json={
        'sheet': 'Sheet1', 'range': 'A1:B3', 'chart_type': 'bar', 'title': 'Sales'
    })
    assert chart.status_code == 200
    assert chart.get_json()['autosaved'] is True
    workbook = client.get(f'/workbooks/{book_id}/save')
    assert workbook.status_code == 200 and workbook.data[:2] == b'PK'
    listed = client.get('/recovery').get_json()['items']
    assert any(item['workbook_id'] == book_id for item in listed)


def test_undo_redo_cell_edit_roundtrip():
    book_id = client.post('/workbooks', json={}).get_json()['workbook_id']
    client.patch(f'/workbooks/{book_id}/cells', json={'sheet': 'Sheet1', 'cells': [{'cell': 'A1', 'value': 'before'}]})
    client.patch(f'/workbooks/{book_id}/cells', json={'sheet': 'Sheet1', 'cells': [{'cell': 'A1', 'value': 'after'}]})
    assert client.get(f'/workbooks/{book_id}/range', query_string={'sheet': 'Sheet1', 'range': 'A1:A1'}).get_json()['rows'][0][0]['value'] == 'after'
    undone = client.post(f'/workbooks/{book_id}/undo')
    assert undone.status_code == 200 and undone.get_json()['undone'] is True
    assert client.get(f'/workbooks/{book_id}/range', query_string={'sheet': 'Sheet1', 'range': 'A1:A1'}).get_json()['rows'][0][0]['value'] == 'before'
    redone = client.post(f'/workbooks/{book_id}/redo')
    assert redone.status_code == 200 and redone.get_json()['redone'] is True
    assert client.get(f'/workbooks/{book_id}/range', query_string={'sheet': 'Sheet1', 'range': 'A1:A1'}).get_json()['rows'][0][0]['value'] == 'after'


def test_undo_reports_empty_history_without_error():
    book_id = client.post('/workbooks', json={}).get_json()['workbook_id']
    response = client.post(f'/workbooks/{book_id}/undo')
    assert response.status_code == 200
    assert response.get_json()['undone'] is False


def test_explicit_recovery_restore_endpoint_reloads_saved_workbook():
    created = client.post('/workbooks', json={'sheet_name': 'SavedBook'}).get_json()
    book_id = created['workbook_id']
    client.patch(f'/workbooks/{book_id}/cells', json={
        'sheet': 'SavedBook', 'cells': [{'cell': 'C3', 'value': 'persisted'}]
    })
    _workbooks.clear()
    restored = client.post(f'/recovery/{book_id}/restore')
    assert restored.status_code == 200
    assert restored.get_json()['restored'] is True
    value = client.get(f'/workbooks/{book_id}/range', query_string={
        'sheet': 'SavedBook', 'range': 'C3:C3'
    }).get_json()['rows'][0][0]['value']
    assert value == 'persisted'


def test_recovery_button_and_last_session_resume_are_in_ui():
    response = client.get('/')
    assert response.status_code == 200
    assert 'recoverBtn' in response.get_data(as_text=True)
    assert 'niceexcel.lastWorkbookId' in response.get_data(as_text=True)


def test_phase8_common_text_and_rounding_formulas():
    from niceexcel_api.formula_engine import calculate_cell
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws['A1'] = '  nice Excel  '
    ws['B1'] = '=UPPER(TRIM(A1))'
    ws['B2'] = '=LEN(A1)'
    ws['B3'] = '=LEFT("NiceExcel",4)'
    ws['B4'] = '=ROUND(12.345,2)'
    ws['B5'] = '=ROUNDUP(12.341,2)'
    ws['B6'] = '=ROUNDDOWN(12.349,2)'
    ws['B7'] = '=IFERROR(1/0,"safe")'
    assert calculate_cell(wb, ws.title, 'B1')['calculated_value'] == 'NICE EXCEL'
    assert calculate_cell(wb, ws.title, 'B2')['calculated_value'] == 14
    assert calculate_cell(wb, ws.title, 'B3')['calculated_value'] == 'Nice'
    assert calculate_cell(wb, ws.title, 'B4')['calculated_value'] == 12.35
    assert calculate_cell(wb, ws.title, 'B5')['calculated_value'] == 12.35
    assert calculate_cell(wb, ws.title, 'B6')['calculated_value'] == 12.34
    assert calculate_cell(wb, ws.title, 'B7')['calculated_value'] == 'safe'


def test_phase8_calculate_workbook_endpoint_evaluates_formulas_and_autosaves():
    created = client.post('/workbooks', json={'sheet_name': 'Calc'}).get_json()
    book_id = created['workbook_id']
    client.patch(f'/workbooks/{book_id}/cells', json={'sheet': 'Calc', 'cells': [
        {'cell': 'A1', 'value': 10}, {'cell': 'A2', 'value': 5},
        {'cell': 'B1', 'value': '=SUM(A1:A2)'}, {'cell': 'B2', 'value': '=IFERROR(1/0,"fallback")'},
        {'cell': 'B3', 'value': '=MISSING(A1)'},
    ]})
    response = client.post(f'/workbooks/{book_id}/calculate')
    assert response.status_code == 200
    data = response.get_json()
    assert data['formulas_evaluated'] == 3
    assert data['errors'] == 1
    values = {item['cell']: item['calculated_value'] for item in data['results']}
    assert values == {'B1': 15, 'B2': 'fallback', 'B3': '#NAME?'}
    assert data['autosaved'] is True
    saved = client.get(f'/workbooks/{book_id}/range', query_string={'sheet': 'Calc', 'range': 'B1:B3'}).get_json()
    assert saved['rows'][0][0]['value'] == '=SUM(A1:A2)'  # formulas stay intact


def test_phase9_cross_sheet_references_and_live_dependency_recalculation():
    created = client.post('/workbooks', json={'sheet_name': 'Sales Data'}).get_json()
    book_id = created['workbook_id']
    assert client.post(f'/workbooks/{book_id}/sheets', json={'name': 'Summary'}).status_code == 200
    client.patch(f'/workbooks/{book_id}/cells', json={'sheet': 'Sales Data', 'cells': [
        {'cell': 'A1', 'value': 10}, {'cell': 'A2', 'value': 20},
        {'cell': 'B1', 'value': '=A1*2'},
    ]})
    client.patch(f'/workbooks/{book_id}/cells', json={'sheet': 'Summary', 'cells': [
        {'cell': 'A1', 'value': "='Sales Data'!A1"},
        {'cell': 'A2', 'value': "=SUM('Sales Data'!A1:A2)"},
        {'cell': 'A3', 'value': "='Sales Data'!B1+5"},
    ]})
    def values():
        return [r[0]['calculated_value'] for r in client.get(
            f'/workbooks/{book_id}/range', query_string={'sheet': 'Summary', 'range': 'A1:A3'}
        ).get_json()['rows']]
    assert values() == [10, 30.0, 25]
    # Changing the source cell automatically changes dependent results on the next refresh.
    client.patch(f'/workbooks/{book_id}/cells', json={'sheet': 'Sales Data', 'cells': [{'cell': 'A1', 'value': 7}]})
    assert values() == [7, 27.0, 19]


def test_phase9_cross_sheet_lookup_and_circular_reference_detection():
    created = client.post('/workbooks', json={'sheet_name': 'Data Sheet'}).get_json()
    book_id = created['workbook_id']
    client.post(f'/workbooks/{book_id}/sheets', json={'name': 'Calc'})
    client.patch(f'/workbooks/{book_id}/cells', json={'sheet': 'Data Sheet', 'cells': [
        {'cell': 'A1', 'value': 'Apple'}, {'cell': 'B1', 'value': 42},
    ]})
    client.patch(f'/workbooks/{book_id}/cells', json={'sheet': 'Calc', 'cells': [
        {'cell': 'A1', 'value': "=VLOOKUP(\"Apple\",'Data Sheet'!A1:B1,2,FALSE)"},
        {'cell': 'B1', 'value': '=C1'}, {'cell': 'C1', 'value': '=B1'},
    ]})
    rows = client.get(f'/workbooks/{book_id}/range', query_string={'sheet': 'Calc', 'range': 'A1:C1'}).get_json()['rows'][0]
    assert rows[0]['calculated_value'] == 42
    assert rows[1]['calculated_value'] == '#CIRCULAR!'
    assert rows[2]['calculated_value'] == '#CIRCULAR!'


def test_formula_preview_evaluates_without_committing_cell_or_history():
    from niceexcel_api.main import _undo_history
    book_id = client.post('/workbooks', json={}).get_json()['workbook_id']
    client.patch(f'/workbooks/{book_id}/cells', json={'sheet': 'Sheet1', 'cells': [
        {'cell': 'A1', 'value': 12}, {'cell': 'B1', 'value': '=A1*2'}
    ]})
    history_before = len(_undo_history.get(book_id, []))
    response = client.post(f'/workbooks/{book_id}/formula-preview', json={
        'sheet': 'Sheet1', 'cell': 'B1', 'formula': '=A1*3'
    })
    assert response.status_code == 200
    assert response.get_json()['calculated_value'] == 36
    assert response.get_json()['preview_only'] is True
    assert client.get(f'/workbooks/{book_id}/range', query_string={'sheet':'Sheet1','range':'B1:B1'}).get_json()['rows'][0][0]['value'] == '=A1*2'
    assert len(_undo_history.get(book_id, [])) == history_before


def test_formula_preview_reports_errors_and_ui_is_wired():
    book_id = client.post('/workbooks', json={}).get_json()['workbook_id']
    response = client.post(f'/workbooks/{book_id}/formula-preview', json={
        'sheet': 'Sheet1', 'cell': 'A1', 'formula': '=1/0'
    })
    assert response.status_code == 200
    assert response.get_json()['formula_error'] == '#DIV/0!' 
    html = client.get('/').get_data(as_text=True)
    assert 'formula-preview' in html
    assert 'scheduleFormulaPreview' in html


def test_dependency_graph_reports_cross_sheet_links_and_dependents():
    book_id = client.post('/workbooks', json={'sheet_name': 'Sales'}).get_json()['workbook_id']
    assert client.post(f'/workbooks/{book_id}/sheets', json={'name': 'Summary'}).status_code == 200
    client.patch(f'/workbooks/{book_id}/cells', json={'sheet': 'Sales', 'cells': [
        {'cell': 'A1', 'value': 12}, {'cell': 'A2', 'value': 8},
    ]})
    client.patch(f'/workbooks/{book_id}/cells', json={'sheet': 'Summary', 'cells': [
        {'cell': 'A1', 'value': "=SUM('Sales'!A1:A2)"}, {'cell': 'B1', 'value': '=A1*2'},
        {'cell': 'C1', 'value': '=IF(1=1,"A1",0)'},
    ]})
    response = client.get(f'/workbooks/{book_id}/dependencies')
    assert response.status_code == 200
    graph = response.get_json()
    assert graph['formula_count'] == 3
    assert graph['cycle_count'] == 0
    a1 = next(n for n in graph['nodes'] if n['sheet'] == 'Summary' and n['cell'] == 'A1')
    assert {'sheet': 'Sales', 'cell': 'A1'} in a1['dependencies']
    assert {'sheet': 'Sales', 'cell': 'A2'} in a1['dependencies']
    b1 = next(n for n in graph['nodes'] if n['sheet'] == 'Summary' and n['cell'] == 'B1')
    assert {'sheet': 'Summary', 'cell': 'B1'} in a1['dependents']
    c1 = next(n for n in graph['nodes'] if n['sheet'] == 'Summary' and n['cell'] == 'C1')
    assert c1['dependencies'] == []

def test_dependency_graph_detects_cycles_and_self_reference():
    book_id = client.post('/workbooks', json={}).get_json()['workbook_id']
    client.patch(f'/workbooks/{book_id}/cells', json={'sheet': 'Sheet1', 'cells': [
        {'cell': 'A1', 'value': '=B1+1'}, {'cell': 'B1', 'value': '=A1+1'},
        {'cell': 'C1', 'value': '=C1'}, {'cell': 'D1', 'value': '=SUM(1,2)'},
    ]})
    graph = client.get(f'/workbooks/{book_id}/dependencies').get_json()
    assert graph['cycle_count'] == 2
    assert graph['formula_count'] == 4
    assert all(n['in_cycle'] for n in graph['nodes'] if n['cell'] in {'A1','B1','C1'})
    assert not next(n for n in graph['nodes'] if n['cell'] == 'D1')['in_cycle']


def test_formula_engine_regression_whole_column_vlookup_and_countif_criteria():
    from niceexcel_api.formula_engine import calculate_cell
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws['A1'] = 'Apple'
    ws['B1'] = 12
    ws['A2'] = 'Pear'
    ws['B2'] = 18
    ws['A3'] = 'Orange'
    ws['B3'] = 25
    ws['D1'] = '=VLOOKUP(A2,A:B,2,FALSE)'
    ws['D2'] = '=COUNTIF(B1:B3,">5")'
    ws['D3'] = '=COUNTIF(B1:B3,">20")'
    ws['D4'] = '=SUMIF(B1:B3,">10",B1:B3)'
    ws['D5'] = '=VLOOKUP("Pear",A:B,2,FALSE)'

    assert calculate_cell(wb, ws.title, 'D1')['calculated_value'] == 18
    assert calculate_cell(wb, ws.title, 'D2')['calculated_value'] == 3
    assert calculate_cell(wb, ws.title, 'D3')['calculated_value'] == 1
    assert calculate_cell(wb, ws.title, 'D4')['calculated_value'] == 55.0
    assert calculate_cell(wb, ws.title, 'D5')['calculated_value'] == 18


def test_retrieve_data_copies_all_matching_rows_to_destination_sheet():
    book_id = client.post('/workbooks', json={}).get_json()['workbook_id']
    client.post(f'/workbooks/{book_id}/sheets', json={'name': 'Sheet2'})
    client.patch(f'/workbooks/{book_id}/cells', json={'sheet': 'Sheet1', 'cells': [
        {'cell': 'A1', 'value': 'ID'}, {'cell': 'B1', 'value': 'Name'}, {'cell': 'C1', 'value': 'Dept'}, {'cell': 'D1', 'value': 'Pay'},
        {'cell': 'A2', 'value': 101}, {'cell': 'B2', 'value': 'Ali'}, {'cell': 'C2', 'value': 'HR'}, {'cell': 'D2', 'value': 3000},
        {'cell': 'A3', 'value': 102}, {'cell': 'B3', 'value': 'Ahmed'}, {'cell': 'C3', 'value': 'IT'}, {'cell': 'D3', 'value': 4500},
        {'cell': 'A4', 'value': 105}, {'cell': 'B4', 'value': 'Sara'}, {'cell': 'C4', 'value': 'Finance'}, {'cell': 'D4', 'value': 4000},
        {'cell': 'A5', 'value': 102}, {'cell': 'B5', 'value': 'Ahmed-2'}, {'cell': 'C5', 'value': 'IT'}, {'cell': 'D5', 'value': 4600},
    ]})
    response = client.post(f'/workbooks/{book_id}/retrieve-data', json={
        'source_sheet':'Sheet1','source_range':'A1:D1000','match_column':'A',
        'lookup_ids':['102','105','108'],'destination_sheet':'Sheet2','destination_cell':'A1'})
    assert response.status_code == 200, response.get_data(as_text=True)
    assert response.get_json()['matched_rows'] == 3
    rows = client.get(f'/workbooks/{book_id}/range', query_string={'sheet':'Sheet2','range':'A1:D3'}).get_json()['rows']
    assert rows[0][0]['value'] == 102 and rows[0][1]['value'] == 'Ahmed'
    assert rows[1][1]['value'] == 'Sara'
    assert rows[2][1]['value'] == 'Ahmed-2'


def test_retrieve_data_rejects_match_column_outside_source_range():
    book_id = client.post('/workbooks', json={}).get_json()['workbook_id']
    client.post(f'/workbooks/{book_id}/sheets', json={'name':'Sheet2'})
    response = client.post(f'/workbooks/{book_id}/retrieve-data', json={
        'source_sheet':'Sheet1','source_range':'A1:D10','match_column':'E',
        'lookup_ids':[102],'destination_sheet':'Sheet2','destination_cell':'A1'})
    assert response.status_code == 422


def test_filter_formula_returns_matching_rows_and_calculate_spills_results():
    book_id = client.post('/workbooks', json={}).get_json()['workbook_id']
    client.post(f'/workbooks/{book_id}/sheets', json={'name':'Sheet2'})
    client.patch(f'/workbooks/{book_id}/cells', json={'sheet':'Sheet1','cells':[
        {'cell':'A1','value':'ID'},{'cell':'B1','value':'Name'},
        {'cell':'A2','value':101},{'cell':'B2','value':'Ali'},
        {'cell':'A3','value':102},{'cell':'B3','value':'Ahmed'},
        {'cell':'A4','value':105},{'cell':'B4','value':'Sara'},
    ]})
    client.patch(f'/workbooks/{book_id}/cells', json={'sheet':'Sheet2','cells':[
        {'cell':'A1','value':102},{'cell':'A2','value':105},
        {'cell':'C1','value':"=FILTER(Sheet1!A2:B4,ISNUMBER(MATCH(Sheet1!A2:A4,Sheet2!A1:A2,0)))"},
    ]})
    calc = client.post(f'/workbooks/{book_id}/calculate')
    assert calc.status_code == 200, calc.get_data(as_text=True)
    rows = client.get(f'/workbooks/{book_id}/range',query_string={'sheet':'Sheet2','range':'C1:D2'}).get_json()['rows']
    assert rows[0][0]['calculated_value'] == [[102, 'Ahmed'], [105, 'Sara']]
    assert rows[0][1]['value'] == 'Ahmed'
    assert rows[1][0]['value'] == 105
    assert rows[1][1]['value'] == 'Sara'


def test_history_batch_undo_redo():
    from niceexcel_api.main import _undo_history, _redo_history
    book_id = client.post('/workbooks', json={}).get_json()['workbook_id']
    client.patch(f'/workbooks/{book_id}/cells', json={
        'sheet': 'Sheet1', 'cells': [{'cell': 'A1', 'value': 'before'}]
    })
    # Make an explicit history point using the same internal recorder as normal edits.
    from niceexcel_api.main import _record_history
    _record_history(book_id)
    client.patch(f'/workbooks/{book_id}/cells', json={
        'sheet': 'Sheet1', 'cells': [{'cell': 'A1', 'value': 'after'}]
    })
    response = client.post(f'/workbooks/{book_id}/history-batch', json={'actions': ['undo', 'redo']})
    assert response.status_code == 200
    assert response.get_json()['changed'] is True
    assert len(response.get_json()['results']) == 2
