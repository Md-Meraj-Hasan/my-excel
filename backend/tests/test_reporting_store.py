import json
from openpyxl import Workbook, load_workbook
from niceexcel_api.reporting_metadata import empty_metadata
from niceexcel_api.reporting_store import (META_SHEET_NAME, get_reporting_metadata, set_reporting_metadata, visible_sheetnames, ensure_excel_tables)


def test_reporting_metadata_is_embedded_and_hidden_across_xlsx_roundtrip():
    wb = Workbook()
    wb.active.title = "Sales"
    wb["Sales"]["A1"] = "Region"
    wb["Sales"]["B1"] = "Sales"
    doc = {"tables": [{"id": "sales_table", "sheet": "Sales", "source_range": "A1:B4", "style": "blue"}]}
    set_reporting_metadata(wb, doc)
    assert wb[META_SHEET_NAME].sheet_state == "veryHidden"
    assert visible_sheetnames(wb) == ["Sales"]
    import io
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    loaded = load_workbook(stream)
    assert get_reporting_metadata(loaded)["tables"][0]["source_range"] == "A1:B4"
    assert loaded[META_SHEET_NAME].sheet_state == "veryHidden"


def test_metadata_roundtrips_when_json_requires_multiple_cells():
    wb = Workbook()
    huge = {"dashboards": [{"id": "dashboard", "name": "D", "components": [
        {"id": f"text_{i}", "type": "text", "x": i, "y": 0, "width": 1, "height": 1, "text": "x" * 120}
        for i in range(300)]}]}
    set_reporting_metadata(wb, huge)
    assert wb[META_SHEET_NAME].max_row > 1
    assert len(get_reporting_metadata(wb)["dashboards"][0]["components"]) == 300


def test_missing_or_corrupt_metadata_is_safe():
    wb = Workbook()
    assert get_reporting_metadata(wb) == empty_metadata()
    ws = wb.create_sheet(META_SHEET_NAME)
    ws["A1"] = "not json"
    assert get_reporting_metadata(wb) == empty_metadata()


def test_table_definitions_rehydrate_native_excel_filter_tables():
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Sales"])
    ws.append(["East", 100])
    ws.append(["West", 200])
    set_reporting_metadata(wb, {"tables": [{
        "id": "sales_table", "name": "Sales", "sheet": "Sales", "source_range": "A1:B3",
        "style": "green", "filter_enabled": True, "total_row": False,
        "excel_table_name": "tbl_sales_table"
    }]})
    assert ensure_excel_tables(wb) == 1
    assert "tbl_sales_table" in ws.tables
    assert ws.tables["tbl_sales_table"].ref == "A1:B3"
    assert ensure_excel_tables(wb) == 0


def test_rehydrated_table_and_hidden_metadata_survive_another_xlsx_roundtrip():
    import io
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory"
    ws.append(["Item", "Qty"])
    ws.append(["Cable", 12])
    set_reporting_metadata(wb, {"tables": [{
        "id": "inventory_table", "name": "Inventory", "sheet": "Inventory",
        "source_range": "A1:B2", "data_range": "A1:B2", "style": "orange",
        "filter_enabled": True, "total_row": False, "excel_table_name": "tbl_inventory_table"
    }]})
    assert ensure_excel_tables(wb) == 1
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    loaded = load_workbook(stream)
    assert "tbl_inventory_table" in loaded["Inventory"].tables
    assert loaded[META_SHEET_NAME].sheet_state == "veryHidden"
    assert get_reporting_metadata(loaded)["tables"][0]["style"] == "orange"
