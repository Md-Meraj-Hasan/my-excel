from openpyxl import Workbook
from niceexcel_api.formula_engine import calculate_cell


def run_formula(formula, values):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for address, value in values.items():
        ws[address] = value
    ws["Z1"] = formula
    return calculate_cell(wb, "Sheet1", "Z1")


def test_xlookup_text_is_case_insensitive():
    result = run_formula('=XLOOKUP("rAvI",A2:A4,B2:B4)', {
        "A2": "Nisha", "B2": "N-01", "A3": "Ravi", "B3": "R-02", "A4": "Omar", "B4": "O-03"
    })
    assert result["calculated_value"] == "R-02"
    assert result["formula_error"] is None


def test_xlookup_matches_numeric_text_and_number():
    result = run_formula("=XLOOKUP(102,A1:A3,B1:B3)", {
        "A1": "101", "B1": "A", "A2": "102", "B2": "B", "A3": "103", "B3": "C"
    })
    assert result["calculated_value"] == "B"


def test_xlookup_if_not_found():
    result = run_formula('=XLOOKUP("missing",A1:A2,B1:B2,"not found")', {
        "A1": "A", "B1": 1, "A2": "B", "B2": 2
    })
    assert result["calculated_value"] == "not found"


def test_xlookup_reverse_search():
    result = run_formula('=XLOOKUP("A",A1:A3,B1:B3,"",0,-1)', {
        "A1": "A", "B1": 1, "A2": "A", "B2": 2, "A3": "A", "B3": 3
    })
    assert result["calculated_value"] == 3


def test_sumifs_and_countifs():
    result = run_formula('=SUMIFS(B2:B4,A2:A4,"East")', {"A2":"East","B2":100,"A3":"West","B3":50,"A4":"East","B4":25})
    assert result["calculated_value"] == 125
    result = run_formula('=COUNTIFS(A2:A4,"East",B2:B4,">50")', {"A2":"East","B2":100,"A3":"West","B3":50,"A4":"East","B4":25})
    assert result["calculated_value"] == 1


def test_averageif_and_averageifs():
    result = run_formula('=AVERAGEIF(A2:A4,"East",B2:B4)', {"A2":"East","B2":100,"A3":"West","B3":50,"A4":"East","B4":25})
    assert result["calculated_value"] == 62.5
    result = run_formula('=AVERAGEIFS(B2:B4,A2:A4,"East",B2:B4,">50")', {"A2":"East","B2":100,"A3":"West","B3":50,"A4":"East","B4":25})
    assert result["calculated_value"] == 100


def test_median_product_and_ifna():
    assert run_formula('=MEDIAN(A1:A3)', {"A1":1,"A2":9,"A3":5})["calculated_value"] == 5
    assert run_formula('=PRODUCT(A1:A3)', {"A1":2,"A2":3,"A3":4})["calculated_value"] == 24
    assert run_formula('=IFNA(XLOOKUP("z",A1:A2,B1:B2),"missing")', {"A1":"a","B1":1,"A2":"b","B2":2})["calculated_value"] == "missing"


def test_date_functions_and_ifna():
    assert run_formula('=DATE(2024,1,1)', {})['calculated_value'] == 45292
    assert run_formula('=YEAR(45292)', {})['calculated_value'] == 2024
    assert run_formula('=MONTH(45292)', {})['calculated_value'] == 1
    assert run_formula('=DAY(45292)', {})['calculated_value'] == 1
    assert run_formula('=IFNA(XLOOKUP("z",A1:A2,B1:B2),"missing")', {'A1':'a','B1':1,'A2':'b','B2':2})['calculated_value'] == 'missing'
