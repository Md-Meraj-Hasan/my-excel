"""Workbook-embedded persistence for NiceExcel reporting definitions."""
from __future__ import annotations

import json
from typing import Any
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from .reporting_metadata import empty_metadata, validate_metadata

META_SHEET_NAME = "_NiceExcelMeta"
META_CHUNK_SIZE = 30000  # below Excel's 32,767-character cell limit


def visible_sheetnames(wb: Workbook) -> list[str]:
    return [name for name in wb.sheetnames if name != META_SHEET_NAME]


def get_reporting_metadata(wb: Workbook) -> dict[str, Any]:
    ws = wb[META_SHEET_NAME] if META_SHEET_NAME in wb.sheetnames else None
    if ws is None:
        return empty_metadata()
    chunks = []
    for row in range(1, ws.max_row + 1):
        value = ws.cell(row, 1).value
        if value is None:
            break
        chunks.append(str(value))
    if not chunks:
        return empty_metadata()
    try:
        return validate_metadata(json.loads("".join(chunks)))
    except (ValueError, TypeError, json.JSONDecodeError):
        # Optional metadata must never prevent opening an otherwise valid workbook.
        return empty_metadata()


def set_reporting_metadata(wb: Workbook, document: Any) -> dict[str, Any]:
    normalized = validate_metadata(document)
    serialized = json.dumps(normalized, ensure_ascii=False, separators=(",", ":"))
    chunks = [serialized[i:i + META_CHUNK_SIZE] for i in range(0, len(serialized), META_CHUNK_SIZE)] or [""]
    ws = wb[META_SHEET_NAME] if META_SHEET_NAME in wb.sheetnames else wb.create_sheet(META_SHEET_NAME)
    ws.delete_rows(1, ws.max_row)
    for row, chunk in enumerate(chunks, start=1):
        ws.cell(row, 1, chunk)
    ws.sheet_state = "veryHidden"
    return normalized


def ensure_excel_tables(wb: Workbook) -> int:
    """Rebuild native XLSX table objects from metadata after a lightweight editor save.

    The Android cell codec preserves cell styling and this hidden metadata sheet, but
    does not serialize OOXML tableParts/relationships itself. Rehydration on the web
    engine restores native Excel filter/table definitions when the workbook is reopened.
    """
    metadata = get_reporting_metadata(wb)
    styles = {
        "blue": "TableStyleMedium2", "green": "TableStyleMedium4",
        "purple": "TableStyleMedium5", "orange": "TableStyleMedium6",
        "teal": "TableStyleMedium9", "gray": "TableStyleMedium1",
        "dark": "TableStyleMedium3", "light": "TableStyleLight9",
    }
    rebuilt = 0
    for definition in metadata["tables"]:
        sheet_name = definition.get("sheet")
        if sheet_name not in wb.sheetnames or sheet_name == META_SHEET_NAME:
            continue
        ws = wb[sheet_name]
        ref = definition.get("source_range")
        if not ref:
            continue
        if any(table.ref == ref for table in ws.tables.values()):
            continue
        name = definition.get("excel_table_name") or ("tbl_" + definition["id"].replace("-", "_"))
        used = {table.displayName.casefold() for sheet in wb.worksheets for table in sheet.tables.values()}
        if name.casefold() in used:
            name = f"{name}_{rebuilt + 1}"
        table = Table(displayName=name, ref=ref)
        style = styles.get(definition.get("style", "blue"), "TableStyleMedium2")
        table.tableStyleInfo = TableStyleInfo(name=style, showFirstColumn=False, showLastColumn=False,
                                              showRowStripes=True, showColumnStripes=False)
        table.totalsRowShown = bool(definition.get("total_row", False))
        if not definition.get("filter_enabled", True):
            table.autoFilter = None
        ws.add_table(table)
        rebuilt += 1
    return rebuilt
